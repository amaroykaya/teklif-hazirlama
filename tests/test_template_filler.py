from datetime import date
from pathlib import Path

import openpyxl

from teklif_hazirlama.application.quote_workflow import QuoteRequest, QuoteWorkflow


def _generate(
    import_path: str,
    teklif_no: str = "RKTSN-2606-26-1-R3",
    ozel_sartlar: list[str] | None = None,
    sartlar: list[str] | None = None,
) -> Path:
    from teklif_hazirlama.core.template_filler import load_default_sartlar

    out = Path("tests/output")
    out.mkdir(exist_ok=True)
    workflow = QuoteWorkflow()
    result = workflow.generate(
        QuoteRequest(
            customer_code="roketsan",
            teklif_no=teklif_no,
            hitap_kisi="Test Kisi",
            hitap_adres_satirlari=[],
            hitap_telefon="",
            hazirlayan="Alican Uzun",
            teslimat="Yurtiçi Kargo",
            teslimat_sekli="Kapı Teslim",
            odeme_sekli="Peşin",
            import_excel_path=import_path,
            ozel_sartlar=ozel_sartlar or [],
            sartlar=sartlar if sartlar is not None else load_default_sartlar(),
            output_dir=str(out),
            tarih=date(2026, 7, 9),
        ),
        create_pdf=False,
    )
    return result.document_path


def _find_sartlar_row(ws):
    for row in range(25, 80):
        value = ws.cell(row, 1).value
        if value and str(value).strip().lower().startswith("şartlar"):
            return row
    return None


def test_teslim_tarihi_row_autofit():
    fixture = Path("tests/fixtures/QBR2525.xlsx")
    if not fixture.exists():
        return

    path = _generate(str(fixture), teklif_no="TEST-TESLIM-HEIGHT")
    wb = openpyxl.load_workbook(path)
    ws = wb["Teklif"]
    assert ws["B23"].alignment.wrap_text is True
    assert (ws.row_dimensions[23].height or 0) > 45
    text = str(ws["B23"].value or "")
    assert "\n" in text
    assert "\n\n" not in text


def test_aciklama_row_autofit_and_wrap():
    from decimal import Decimal
    from datetime import date

    from teklif_hazirlama.core.models import QuoteDocument, QuoteHeader, QuoteLine
    from teklif_hazirlama.core.template_filler import TemplateFiller, load_default_sartlar
    from teklif_hazirlama.infrastructure.storage import CustomerRepository

    cust = CustomerRepository().load("roketsan")
    long = (
        "VIDEO ENCODER DVI TO SDI CONVERTER UNIT WITH EXTENDED DESCRIPTION "
        "/ 00239277\n( K ve P pozisyonlari fiyata dahil degildir. )"
    )
    lines = [
        QuoteLine(
            row_number=1,
            adet=1,
            ants_is_urun_kodu="ANT-1",
            aciklama=long,
            birim_fiyat=Decimal("10"),
            toplam_fiyat=Decimal("10"),
        )
    ]
    header = QuoteHeader(
        teklif_no="TEST-ACIKLAMA",
        tarih=date(2026, 7, 9),
        hitap_kisi="Test",
        hitap_adres_satirlari=[],
        hitap_telefon="",
        hazirlayan="Alican Uzun",
        teslimat="Yurtiçi Kargo",
        teslimat_sekli="Kapı Teslim",
        odeme_sekli="Peşin",
        istek_no="QR1",
    )
    doc = QuoteDocument(
        customer=cust,
        header=header,
        lines=lines,
        ara_toplam=Decimal("10"),
        kdv=Decimal("2"),
        toplam=Decimal("12"),
        teslim_tarihi_text="x",
        sartlar=load_default_sartlar(),
    )
    out = Path("tests/output")
    out.mkdir(exist_ok=True)
    path = out / "test-aciklama-autofit.xlsx"
    TemplateFiller().fill(doc, path)
    ws = openpyxl.load_workbook(path)["Teklif"]
    assert ws["C27"].alignment.wrap_text is True
    assert ws["C27"].alignment.vertical == "top"
    assert (ws.row_dimensions[27].height or 0) >= 45


def test_teklif_no_widens_h_column():
    fixture = Path("tests/fixtures/QBR2525.xlsx")
    if not fixture.exists():
        return

    long_no = "RKTSN-2606-26-1-R3-COK-UZUN-TEKLIF"
    path = _generate(str(fixture), teklif_no=long_no)
    wb = openpyxl.load_workbook(path)
    ws = wb["Teklif"]
    assert ws["H9"].value == long_no
    assert (ws.column_dimensions["H"].width or 0) > 13
    assert ws["H9"].alignment.shrink_to_fit is True


def test_teslimat_sekli_column_widens():
    from datetime import date

    from teklif_hazirlama.application.quote_workflow import QuoteRequest, QuoteWorkflow
    from teklif_hazirlama.core.template_filler import load_default_sartlar

    fixture = Path("tests/fixtures/QBR2525.xlsx")
    if not fixture.exists():
        return

    out = Path("tests/output")
    out.mkdir(exist_ok=True)
    long_sekli = "Kapı Teslim + Montaj Dahil Ekstra"
    result = QuoteWorkflow().generate(
        QuoteRequest(
            customer_code="roketsan",
            teklif_no="TEST-TESLIMAT-SEKLI",
            hitap_kisi="Test",
            hitap_adres_satirlari=[],
            hitap_telefon="",
            hazirlayan="Alican Uzun",
            teslimat="Yurtiçi Kargo",
            teslimat_sekli=long_sekli,
            odeme_sekli="Peşin",
            import_excel_path=str(fixture),
            ozel_sartlar=[],
            sartlar=load_default_sartlar(),
            output_dir=str(out),
            tarih=date(2026, 7, 9),
        ),
        create_pdf=False,
    )
    ws = openpyxl.load_workbook(result.document_path)["Teklif"]
    assert ws["F23"].value == long_sekli
    assert ws["F23"].alignment.wrap_text is True
    assert (ws.column_dimensions["F"].width or 0) > 13


def test_print_area_covers_sartlar_bottom():
    real_file = Path(r"C:\Users\Asus.DESKTOP-9F6EQVL\Desktop\KLDHDh\QR075313.xlsx")
    if not real_file.exists():
        return

    path = _generate(str(real_file), teklif_no="TEST-PRINT-AREA", ozel_sartlar=["Ozel"])
    wb = openpyxl.load_workbook(path)
    ws = wb["Teklif"]
    sartlar_row = _find_sartlar_row(ws)
    assert sartlar_row is not None
    last_content = sartlar_row
    for row in range(sartlar_row, sartlar_row + 30):
        value = ws.cell(row, 1).value
        if value and str(value).strip():
            last_content = row
    area = (ws.print_area or "").replace("$", "")
    # print area şartlar son satırını kapsar
    assert "A1:H" in area
    end = int(area.split("H")[-1])
    assert end >= last_content


def test_sartlar_written_below_header_for_two_lines():
    fixture = Path("tests/fixtures/QBR2525.xlsx")
    if not fixture.exists():
        return

    path = _generate(str(fixture), teklif_no="TEST-2LINES")
    ws = openpyxl.load_workbook(path, data_only=True)["Teklif"]
    sartlar_row = _find_sartlar_row(ws)
    assert sartlar_row == 35
    # Teşekkür metni Şartlar'dan önce durur
    assert "teşekkür" in str(ws.cell(33, 1).value or "").lower()
    # Şartlar: hemen altı = ürün kodları; diğer şart metinleri daha aşağıda
    assert "ANT-" in str(ws.cell(sartlar_row + 1, 1).value or "")
    found_fixed = False
    for row in range(sartlar_row + 2, sartlar_row + 20):
        value = str(ws.cell(row, 1).value or "")
        if value.startswith("Teklif edilen"):
            found_fixed = True
            break
    assert found_fixed


def test_codes_and_ozel_insert_two_rows_under_sartlar():
    fixture = Path("tests/fixtures/QBR2525.xlsx")
    if not fixture.exists():
        return

    path = _generate(
        str(fixture),
        teklif_no="TEST-2-ROWS",
        ozel_sartlar=["Ozel sart A"],
    )
    ws = openpyxl.load_workbook(path, data_only=True)["Teklif"]
    sartlar_row = _find_sartlar_row(ws)
    assert sartlar_row is not None
    # Hemen altı: kodlar, sonra özel şart, sonra diğer yazılar
    assert "ANT-" in str(ws.cell(sartlar_row + 1, 1).value or "")
    assert "Ozel sart A" in str(ws.cell(sartlar_row + 2, 1).value or "")
    assert str(ws.cell(sartlar_row + 3, 1).value or "").startswith("Teklif edilen")


def test_sartlar_and_ozel_sart_below_header_for_many_lines():
    real_file = Path(r"C:\Users\Asus.DESKTOP-9F6EQVL\Desktop\KLDHDh\QR075313.xlsx")
    if not real_file.exists():
        return

    path = _generate(
        str(real_file),
        teklif_no="TEST-8LINES",
        ozel_sartlar=["Manuel ozel sart 1", "Manuel ozel sart 2"],
    )
    ws = openpyxl.load_workbook(path, data_only=True)["Teklif"]
    sartlar_row = _find_sartlar_row(ws)
    assert sartlar_row is not None
    assert "ANT-DVI-ENC-623F" in str(ws.cell(sartlar_row + 1, 1).value or "")
    assert "Manuel ozel sart 1" in str(ws.cell(sartlar_row + 2, 1).value or "")
    assert str(ws.cell(sartlar_row + 3, 1).value or "").startswith("Teklif edilen")


def test_editable_sartlar_can_be_removed():
    fixture = Path("tests/fixtures/QBR2525.xlsx")
    if not fixture.exists():
        return
    path = _generate(str(fixture), teklif_no="TEST-SART-EMPTY", sartlar=[])
    ws = openpyxl.load_workbook(path, data_only=True)["Teklif"]
    sartlar_row = _find_sartlar_row(ws)
    assert sartlar_row is not None
    first = str(ws.cell(sartlar_row + 1, 1).value or "")
    assert "ANT-" in first
    assert not first.startswith("Teklif edilen")


def test_revize_template_has_indirim_columns():
    from datetime import date
    from decimal import Decimal

    from teklif_hazirlama.core.models import QuoteDocument, QuoteHeader, QuoteLine
    from teklif_hazirlama.core.template_filler import TemplateFiller, load_default_sartlar
    from teklif_hazirlama.infrastructure.storage import CustomerRepository
    from teklif_hazirlama.paths import REVIZE_TEMPLATE_FILE

    assert REVIZE_TEMPLATE_FILE.exists()
    cust = CustomerRepository().load("roketsan")
    lines = [
        QuoteLine(
            row_number=1,
            adet=2,
            ants_is_urun_kodu="ANT-1",
            aciklama="urun",
            birim_fiyat=Decimal("100"),
            toplam_fiyat=Decimal("200"),
        )
    ]
    header = QuoteHeader(
        teklif_no="RKTSN-2607-1-R2",
        tarih=date(2026, 7, 12),
        hitap_kisi="Test",
        hitap_adres_satirlari=[],
        hitap_telefon="",
        hazirlayan="Alican Uzun",
        teslimat="Yurtiçi Kargo",
        teslimat_sekli="Kapı Teslim",
        odeme_sekli="Peşin",
        istek_no="QR1",
        revizyon="2",
    )
    doc = QuoteDocument(
        customer=cust,
        header=header,
        lines=lines,
        ara_toplam=Decimal("200"),
        kdv=Decimal("40"),
        toplam=Decimal("240"),
        teslim_tarihi_text="x",
        sartlar=load_default_sartlar(revise=True),
        ozel_sartlar=["Ozel"],
    )
    out = Path("tests/output")
    out.mkdir(exist_ok=True)
    path = out / "test-revize.xlsx"
    TemplateFiller(revise=True).fill(doc, path)
    ws = openpyxl.load_workbook(path)["Revize Teklif"]
    assert ws["E26"].value == "Birim Fiyat"
    assert "ndirim" in str(ws["F26"].value or "").lower()
    assert ws["E27"].value == 100.0
    assert ws["G27"].value == 100.0
    assert str(ws["F27"].value).replace(" ", "") == "=1-(G27/E27)"
    assert "ANT-" in str(ws.cell(_find_sartlar_row(ws) + 1, 1).value or "")


def test_load_default_sartlar_from_template():
    from teklif_hazirlama.core.template_filler import load_default_sartlar

    lines = load_default_sartlar()
    assert len(lines) >= 5
    assert lines[0].startswith("Teklif edilen")


def test_no_stray_merge_over_sartlar_text():
    real_file = Path(r"C:\Users\Asus.DESKTOP-9F6EQVL\Desktop\KLDHDh\QR075313.xlsx")
    if not real_file.exists():
        return

    path = _generate(str(real_file), teklif_no="TEST-NO-MERGE", ozel_sartlar=["Ozel"])
    wb = openpyxl.load_workbook(path)
    ws = wb["Teklif"]
    sartlar_row = _find_sartlar_row(ws)
    assert sartlar_row is not None
    stray = [
        str(m)
        for m in ws.merged_cells.ranges
        if m.min_row >= sartlar_row and m.min_col <= 6 and m.max_col >= 2
    ]
    assert stray == []
    # Truncated-looking şartlar lines must still hold full text in A
    for offset in (4, 5):
        value = str(ws.cell(sartlar_row + offset, 1).value or "")
        assert len(value) > 40
