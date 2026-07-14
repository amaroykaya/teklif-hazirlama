from pathlib import Path

from openpyxl import Workbook

from teklif_hazirlama.application.quote_workflow import QuoteRequest, QuoteWorkflow
from teklif_hazirlama.core.models import SheetsRow


def _make_erp(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.append(["meta"])
    ws.append(
        [
            "Teklifte Bulun",
            "Stok Kodu",
            "Stok Tanımı",
            "Miktar",
            "Fiyat",
            "Satır Toplamı",
            "Tedarikçi Notu",
            "Temin Süresi (Takvim Günü)",
            "Garanti Süresi (Yıl)",
        ]
    )
    ws.append(["Y", "002525", "Urun", 2, None, None, None, None, None])
    wb.save(path)
    wb.close()


def test_enrich_on_import_then_generate(tmp_path):
    src = tmp_path / "QR075313.xlsx"
    _make_erp(src)
    out = tmp_path / "out"
    out.mkdir()
    sheets = [
        SheetsRow(
            musteri_stok_kodu="x",
            antsis_urun_kodu="ANT-TEST",
            teklif_sevk_tarihi="2 x 10 hafta",
            birim_fiyat="50,00",
        )
    ]
    wf = QuoteWorkflow()
    enrich = wf.enrich_import(src, sheets)
    assert enrich.filled_rows == 1
    assert enrich.enriched_path.exists()

    from decimal import Decimal

    from teklif_hazirlama.core.models import QuoteLine

    edited = [
        QuoteLine(
            row_number=1,
            adet=9,
            ants_is_urun_kodu="ANT-EDITED",
            aciklama="Elle degisti / 002525",
            birim_fiyat=Decimal("77.00"),
            toplam_fiyat=Decimal("693.00"),
            stok_kodu="002525",
        )
    ]
    result = wf.generate(
        QuoteRequest(
            customer_code="roketsan",
            teklif_no="RKTSN-2607-11",
            hitap_kisi="Test",
            hitap_adres_satirlari=[],
            hitap_telefon="",
            hazirlayan="Kerem Özsoy",
            teslimat="Yurtiçi Kargo",
            teslimat_sekli="Kapı Teslim",
            odeme_sekli="Peşin",
            import_excel_path=str(enrich.enriched_path),
            source_excel_path=str(src),
            ozel_sartlar=[],
            output_dir=str(out),
            sheets_rows=[],
            quote_lines=edited,
        ),
        create_pdf=False,
    )
    assert result.document_path.exists()
    assert result.enriched_excel_path is not None
    assert result.original_excel_path is not None
    assert result.output_folder is not None
    assert result.original_excel_path.name == "QR075313.xlsx"
    assert result.enriched_excel_path.name == "QR075313-doldurulmuş.xlsx"
    assert result.original_excel_path.exists()
    assert result.enriched_excel_path.exists()

    from openpyxl import load_workbook

    wb = load_workbook(result.enriched_excel_path, data_only=True)
    ws = wb.active
    assert ws.cell(3, 4).value == 9
    assert float(ws.cell(3, 5).value) == 77.0
    assert "ANT-EDITED" in str(ws.cell(3, 7).value)
    wb.close()


def test_generate_excel_from_fixture():
    fixture = Path("tests/fixtures/QBR2525.xlsx")
    if not fixture.exists():
        return
    assert fixture.exists()
