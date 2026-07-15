from openpyxl import Workbook, load_workbook

from teklif_hazirlama.core.excel_enricher import ExcelEnricher
from teklif_hazirlama.core.import_parser import ImportParser
from teklif_hazirlama.core.models import SheetsRow
from teklif_hazirlama.core.tedarikci_notu_parser import parse_tedarikci_notu


def _make_erp_workbook(path):
    wb = Workbook()
    ws = wb.active
    ws.append(["meta"])
    ws.append(
        [
            "Sıra",
            "Teklifte Bulun",
            "Stok Kodu",
            "Stok Tanımı",
            "Miktar",
            "Fiyat",
            "Satır Toplamı",
            "Tedarikçi Notu",
            "Temin Süresi (Takvim Günü)",
            "Garanti Süresi (Yıl)",
        ]
    )
    ws.append(
        [
            1,
            "Y",
            "002525",
            "Urun A",
            10,
            None,
            None,
            None,
            None,
            None,
        ]
    )
    ws.append(
        [
            2,
            "Y",
            "999999",
            "Only Excel",
            1,
            1,
            1,
            "Teslim Süreleri : OLD - 1 adet T0+1 Hafta T0: Sipariş Onay Tarihi",
            7,
            1,
        ]
    )
    wb.save(path)
    wb.close()


def test_enrich_fills_by_row_order(tmp_path):
    src = tmp_path / "QR075313.xlsx"
    _make_erp_workbook(src)
    sheets = [
        SheetsRow(
            musteri_stok_kodu="herhangi",
            antsis_urun_kodu="ANT-DVI-ENC-hfkaha",
            teklif_sevk_tarihi="5 x 10 hafta\n10 x 35 hafta",
            birim_fiyat="100,50",
            kalemdeki_ozel_sartlar="K ve P dahil değil.",
        ),
        SheetsRow(
            musteri_stok_kodu="na",
            antsis_urun_kodu="na",
            teklif_sevk_tarihi="na",
            birim_fiyat="na",
        ),
    ]
    out = tmp_path / "out"
    result = ExcelEnricher().enrich(src, sheets, output_dir=out)
    assert result.matched == 2

    wb = load_workbook(result.enriched_path, data_only=True)
    ws = wb.active
    assert ws.cell(3, 6).value == 100.5  # 1. satır Fiyat
    assert ws.cell(3, 9).value == 250
    assert ws.cell(3, 10).value == 1.0
    assert "Teslim süresi : ANT-DVI-ENC-hfkaha" in ws.cell(3, 8).value
    # 2. satır na
    assert ws.cell(4, 6).value == "na"
    assert ws.cell(4, 8).value == "na"
    assert ws.cell(4, 9).value == "na"
    wb.close()

    lines, errors = ImportParser().parse(result.enriched_path)
    assert errors == []
    assert lines[0].ants_is_urun_kodu == "ANT-DVI-ENC-hfkaha"
    assert lines[1].ants_is_urun_kodu == "na"


def test_enrich_skips_n_rows_keeps_them_untouched(tmp_path):
    """CAP-1: N satırı eşlemeyi kaydırmaz; kopyada orijinal değerler kalır."""
    src = tmp_path / "mix.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["meta"])
    ws.append(
        [
            "Sıra",
            "Teklifte Bulun",
            "Stok Kodu",
            "Stok Tanımı",
            "Miktar",
            "Fiyat",
            "Satır Toplamı",
            "Tedarikçi Notu",
            "Temin Süresi (Takvim Günü)",
            "Garanti Süresi (Yıl)",
        ]
    )
    # Y → Sheets[0]
    ws.append([1, "Y", "111", "A", 1, None, None, None, None, None])
    # N → dokunulmaz (eski fiyat 999)
    ws.append([2, "N", "222", "B", 1, 999, 999, "OLD-N", 99, 5])
    # Y → Sheets[1]
    ws.append([3, "Y", "333", "C", 1, None, None, None, None, None])
    wb.save(src)
    wb.close()

    sheets = [
        SheetsRow(
            musteri_stok_kodu="a",
            antsis_urun_kodu="ANT-A",
            teklif_sevk_tarihi="1 x 10 hafta",
            birim_fiyat="10,00",
        ),
        SheetsRow(
            musteri_stok_kodu="c",
            antsis_urun_kodu="ANT-C",
            teklif_sevk_tarihi="2 x 20 hafta",
            birim_fiyat="20,00",
        ),
    ]
    result = ExcelEnricher().enrich(src, sheets, output_dir=tmp_path / "o")
    assert result.matched == 2

    out = load_workbook(result.enriched_path, data_only=True).active
    assert out.cell(3, 6).value == 10.0  # Y1 fiyat
    assert out.cell(3, 10).value == 1.0
    assert out.cell(4, 6).value == 999  # N dokunulmadı
    assert out.cell(4, 8).value == "OLD-N"
    assert out.cell(4, 9).value == 99
    assert out.cell(4, 10).value == 5
    assert out.cell(5, 6).value == 20.0  # Y2 = Sheets[1], N atlandı
    assert "ANT-C" in str(out.cell(5, 8).value or "")


def test_enrich_does_not_mutate_source_import(tmp_path):
    """CAP-2: kaynak import dosyası değişmez."""
    import hashlib

    src = tmp_path / "QR_SRC.xlsx"
    _make_erp_workbook(src)
    before = hashlib.sha256(src.read_bytes()).hexdigest()
    sheets = [
        SheetsRow(
            musteri_stok_kodu="x",
            antsis_urun_kodu="ANT-X",
            teklif_sevk_tarihi="1 x 10 hafta",
            birim_fiyat="9,00",
        )
    ]
    result = ExcelEnricher().enrich(src, sheets, output_dir=tmp_path / "o")
    after = hashlib.sha256(src.read_bytes()).hexdigest()
    assert before == after
    assert result.enriched_path.resolve() != src.resolve()
    assert result.enriched_path.exists()


def test_enrich_extra_excel_rows_get_na(tmp_path):
    src = tmp_path / "QR1.xlsx"
    _make_erp_workbook(src)
    sheets = [
        SheetsRow(
            musteri_stok_kodu="x",
            antsis_urun_kodu="ANT-X",
            teklif_sevk_tarihi="1 x 10 hafta",
            birim_fiyat="9,00",
        )
    ]
    result = ExcelEnricher().enrich(src, sheets, output_dir=tmp_path / "o")
    assert result.matched == 1
    assert any("fazla Y" in w for w in result.warnings)
    wb = load_workbook(result.enriched_path, data_only=True)
    assert wb.active.cell(4, 6).value == "na"
    wb.close()


def test_enrich_n_row_does_not_get_next_y_sheet_values(tmp_path):
    """
    Gerçek QR075313 deseni: Y…Y, N, Y
    Eski hata: N satırı bir alttaki Y'nin Sheets değerlerini alırdı.
    """
    src = tmp_path / "QR075313-pattern.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["meta"])
    ws.append(
        [
            "Satır No",
            "Teklifte Bulun",
            "Stok Kodu",
            "Stok Tanımı",
            "Miktar",
            "Fiyat",
            "Satır Toplamı",
            "Tedarikçi Notu",
            "Temin Süresi (Takvim Günü)",
            "Garanti Süresi (Yıl)",
        ]
    )
    # 5 Y + 1 N + 1 Y
    for i, (tb, stok) in enumerate(
        [
            ("Y", "111"),
            ("Y", "222"),
            ("Y", "333"),
            ("Y", "444"),
            ("Y", "555"),
            ("N", "00236333"),
            ("Y", "EL5079"),
        ],
        start=1,
    ):
        fiyat = 777 if tb == "N" else None
        temin = 88 if tb == "N" else None
        notu = "KEEP-N-USS" if tb == "N" else None
        garanti = "" if tb == "N" else None
        ws.append([i, tb, stok, f"U{i}", 1, fiyat, None, notu, temin, garanti])
    wb.save(src)
    wb.close()

    sheets = [
        SheetsRow(
            musteri_stok_kodu=str(i),
            antsis_urun_kodu=f"ANT-Y{i}",
            teklif_sevk_tarihi=f"1 x {10 + i} hafta",
            birim_fiyat=f"{200 + i},00",
        )
        for i in range(6)
    ]
    # Son Sheets = ALTTAKI Y (EL5079) icin — N buna yazilmamali
    sheets[-1] = SheetsRow(
        musteri_stok_kodu="el",
        antsis_urun_kodu="ANT-0080",
        teklif_sevk_tarihi="1 x 12 hafta",
        birim_fiyat="1600,00",
    )

    result = ExcelEnricher().enrich(src, sheets, output_dir=tmp_path / "o")
    assert result.matched == 6
    out = load_workbook(result.enriched_path, data_only=True).active
    # Excel R8 = N (baslik+meta sonrasi: 3+5=8)
    assert out.cell(8, 2).value == "N"
    assert out.cell(8, 6).value == 777
    assert out.cell(8, 9).value == 88
    assert out.cell(8, 8).value == "KEEP-N-USS"
    # Alttaki Y dogru Sheets'i alir
    assert out.cell(9, 2).value == "Y"
    assert out.cell(9, 6).value == 1600.0
    assert "ANT-0080" in str(out.cell(9, 8).value or "")


def test_parse_tedarikci_notu_accepts_singular_sure():
    text = (
        "Teslim süresi : ANT-DVI-ENC-hfkaha - 5 adet T0+10 Hafta, 10 adet T0+35 Hafta "
        "T0: Sipariş Onay Tarihi K ve P pozisyonları fiyata dahil değildir."
    )
    result = parse_tedarikci_notu(text)
    assert result.urun_kodu == "ANT-DVI-ENC-hfkaha"
    assert "5 adet T0+10 Hafta" in result.teslim_suresi_parcasi
    assert "K ve P" in result.aciklama_eki
