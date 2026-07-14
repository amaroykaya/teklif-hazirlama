from pathlib import Path

from openpyxl import Workbook

from teklif_hazirlama.core.import_parser import ImportParseError, ImportParser
from teklif_hazirlama.infrastructure.excel_import_reader import (
    normalize_header,
    resolve_field_columns,
)


def test_normalize_header_turkish():
    assert normalize_header("Stok Tanımı") == "stok tanimi"
    assert normalize_header("  Tedarikçi Notu ") == "tedarikci notu"


def test_resolve_columns_by_name_ignores_letter_order():
    # Sütunlar kasıtlı olarak klasik sıradan farklı
    headers = [
        "X",
        "Miktar",
        "Teklifte Bulun",
        "Fiyat",
        "Stok Kodu",
        "Stok Tanımı",
        "Satır Toplamı",
        "Tedarikçi Notu",
        "Termin Tarihi",
        "Satır Toplamı EUR",
    ]
    cols = resolve_field_columns(headers)
    assert cols["miktar"] == 1
    assert cols["teklifte_bulun"] == 2
    assert cols["fiyat"] == 3
    assert cols["stok_kodu"] == 4
    assert cols["stok_tanimi"] == 5
    assert cols["satir_toplami"] == 6  # EUR olan değil
    assert cols["tedarikci_notu"] == 7
    assert cols["termin_tarihi"] == 8


def test_parse_shifted_columns_workbook(tmp_path):
    path = tmp_path / "shifted.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["meta"])
    ws.append(
        [
            "Sıra",
            "Miktar",
            "Teklifte Bulun",
            "Fiyat",
            "Stok Kodu",
            "Stok Tanımı",
            "Satır Toplamı",
            "Tedarikçi Notu",
            "Termin Tarihi",
        ]
    )
    ws.append(
        [
            1,
            2,
            "Y",
            10,
            "001",
            "Urun A",
            20,
            "Teslim Süreleri : ANT-TEST - 2 adet T0+1 Hafta T0: Sipariş Onay Tarihi",
            "2026-06-01",
        ]
    )
    wb.save(path)

    lines, errors = ImportParser().parse(path)
    assert errors == []
    assert len(lines) == 1
    assert lines[0].adet == 2
    assert lines[0].stok_aciklama == "Urun A / 001"
    assert lines[0].birim_fiyat == 10
    assert lines[0].toplam_fiyat == 20
    assert lines[0].ants_is_urun_kodu == "ANT-TEST"
    assert lines[0].termin_tarihi is not None


def test_missing_required_column_raises(tmp_path):
    path = tmp_path / "missing.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["meta"])
    ws.append(["Teklifte Bulun", "Stok Kodu", "Miktar"])
    ws.append(["Y", "001", 1])
    wb.save(path)

    try:
        ImportParser().parse(path)
        assert False, "expected ImportParseError"
    except ImportParseError as exc:
        assert "eksik sütun" in str(exc).lower()


def test_parse_istek_without_tedarikci_notu(tmp_path):
    path = tmp_path / "istek.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["meta"])
    ws.append(
        [
            "Teklifte Bulun",
            "Stok Kodu",
            "Stok Tanımı",
            "Miktar",
            "Termin Tarihi",
        ]
    )
    ws.append(["Y", "001", "Urun A", 5, "2026-06-01"])
    ws.append(["Y", "002", "Urun B", 3, "2026-07-01"])
    wb.save(path)

    lines, errors = ImportParser().parse(path, mode="istek")
    assert errors == []
    assert len(lines) == 2
    assert lines[0].stok_aciklama == "Urun A / 001"
    assert lines[0].adet == 5
    assert lines[0].ants_is_urun_kodu == ""
    assert lines[0].termin_tarihi is not None
    real = Path(r"C:\Users\Asus.DESKTOP-9F6EQVL\Desktop\KLDHDh\QR075313.xlsx")
    if not real.exists():
        return
    lines, errors = ImportParser().parse(real)
    assert len(lines) >= 1
    assert lines[0].ants_is_urun_kodu == "ANT-DVI-ENC-623F"
    assert lines[0].adet == 29
    assert lines[0].toplam_fiyat > 0
    assert lines[0].termin_tarihi is not None
