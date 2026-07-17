from decimal import Decimal

from openpyxl import Workbook, load_workbook

from teklif_hazirlama.core.excel_ui_sync import sync_quote_lines_to_import_excel
from teklif_hazirlama.core.models import QuoteLine


def _make_erp(path):
    wb = Workbook()
    ws = wb.active
    ws.append(["meta"])
    ws.append(
        [
            "Teklifte Bulun",
            "Stok Kodu",
            "Stok Tanımı",
            "Miktar",
            "Fiyat",
            "Satır Toplamı",
            "Tedarikçi Notu",
            "Temin Süresi (Takvim Günü)",
            "Garanti Süresi (Yıl)",
        ]
    )
    ws.append(
        [
            "Y",
            "002525",
            "Urun A",
            10,
            100.0,
            1000.0,
            "Teslim süresi : OLD-CODE - 10 adet T0+5 Hafta T0: Sipariş Onay Tarihi",
            35,
            1,
        ]
    )
    ws.append(
        [
            "Y",
            "999999",
            "Urun B",
            1,
            50.0,
            50.0,
            "Teslim süresi : B-CODE - 1 adet T0+2 Hafta T0: Sipariş Onay Tarihi",
            14,
            1,
        ]
    )
    wb.save(path)
    wb.close()


def test_sync_updates_existing_row(tmp_path):
    path = tmp_path / "filled.xlsx"
    _make_erp(path)
    lines = [
        QuoteLine(
            row_number=1,
            adet=7,
            ants_is_urun_kodu="NEW-CODE",
            aciklama="Yeni Aciklama / 002525",
            birim_fiyat=Decimal("12.50"),
            toplam_fiyat=Decimal("87.50"),
            stok_kodu="002525",
        ),
        QuoteLine(
            row_number=2,
            adet=1,
            ants_is_urun_kodu="B-CODE",
            aciklama="Urun B / 999999",
            birim_fiyat=Decimal("50"),
            toplam_fiyat=Decimal("50"),
            stok_kodu="999999",
        ),
    ]
    warnings = sync_quote_lines_to_import_excel(path, lines)
    assert any("güncellendi" in w for w in warnings)

    wb = load_workbook(path, data_only=True)
    ws = wb.active
    assert ws.cell(3, 4).value == 7  # Miktar
    assert float(ws.cell(3, 5).value) == 12.5  # Fiyat
    assert float(ws.cell(3, 6).value) == 87.5  # Satır Toplamı
    assert ws.cell(3, 3).value == "Yeni Aciklama"  # Stok Tanımı
    assert "NEW-CODE" in str(ws.cell(3, 7).value)
    assert "10 adet T0+5 Hafta" in str(ws.cell(3, 7).value)
    wb.close()


def test_sync_adds_and_deletes_rows(tmp_path):
    path = tmp_path / "filled.xlsx"
    _make_erp(path)

    # Tek satıra düşür → 2. satır silinmeli
    one = [
        QuoteLine(
            row_number=1,
            adet=3,
            ants_is_urun_kodu="ONLY",
            aciklama="Tek",
            birim_fiyat=Decimal("1"),
            toplam_fiyat=Decimal("3"),
            stok_kodu="002525",
        )
    ]
    sync_quote_lines_to_import_excel(path, one)
    wb = load_workbook(path)
    ws = wb.active
    assert ws.max_row == 3
    assert ws.cell(3, 4).value == 3
    wb.close()

    # İki satır daha ekle
    three = one + [
        QuoteLine(
            row_number=2,
            adet=2,
            ants_is_urun_kodu="A2",
            aciklama="Iki",
            birim_fiyat=Decimal("4"),
            toplam_fiyat=Decimal("8"),
            stok_kodu="",
        ),
        QuoteLine(
            row_number=3,
            adet=5,
            ants_is_urun_kodu="A3",
            aciklama="Uc",
            birim_fiyat=Decimal("2"),
            toplam_fiyat=Decimal("10"),
            stok_kodu="S3",
        ),
    ]
    warnings = sync_quote_lines_to_import_excel(path, three)
    assert any("eklendi" in w for w in warnings)
    wb = load_workbook(path)
    ws = wb.active
    assert ws.cell(4, 1).value == "Y"
    assert ws.cell(4, 4).value == 2
    assert ws.cell(5, 2).value == "S3"
    assert ws.cell(5, 4).value == 5
    wb.close()


def test_sync_updates_temin_from_teslim_tarihi(tmp_path):
    path = tmp_path / "filled.xlsx"
    _make_erp(path)
    lines = [
        QuoteLine(
            row_number=1,
            adet=5,
            ants_is_urun_kodu="ANT5307",
            aciklama="Urun",
            birim_fiyat=Decimal("10"),
            toplam_fiyat=Decimal("50"),
            stok_kodu="002525",
            # Eski import parçası — UI metni öncelikli olmalı
            teslim_suresi_parcasi="ANT5307 - 5 adet T0+23 Hafta",
        ),
        QuoteLine(
            row_number=2,
            adet=1,
            ants_is_urun_kodu="B-CODE",
            aciklama="Urun B",
            birim_fiyat=Decimal("50"),
            toplam_fiyat=Decimal("50"),
            stok_kodu="999999",
        ),
    ]
    teslim = (
        "ANT5307 - 2 adet T0+23 Hafta 3 adet T0+30 Hafta\n"
        "T0: Sipariş Onay Tarihi"
    )
    warnings = sync_quote_lines_to_import_excel(
        path, lines, teslim_tarihi_text=teslim
    )
    assert any("Temin Süresi" in w for w in warnings)
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    # Temin: max 30 hafta → 210 → bir üst onluk 220 (eski 23 hafta=170 olmamalı)
    assert ws.cell(3, 8).value == 220
    notu = str(ws.cell(3, 7).value)
    assert "ANT5307" in notu
    assert "2 adet T0+23 Hafta" in notu
    assert "3 adet T0+30 Hafta" in notu
    wb.close()


def test_sync_temin_with_code_mismatch_still_uses_ui_text(tmp_path):
    """Tablo kodu ile Teslim satır kodu birebir aynı olmasa da UI metni kullanılsın."""
    path = tmp_path / "filled.xlsx"
    _make_erp(path)
    lines = [
        QuoteLine(
            row_number=1,
            adet=5,
            ants_is_urun_kodu="ANT-5307-FULL",
            aciklama="Urun",
            birim_fiyat=Decimal("10"),
            toplam_fiyat=Decimal("50"),
            stok_kodu="002525",
            teslim_suresi_parcasi="5 adet T0+10 Hafta",
        ),
        QuoteLine(
            row_number=2,
            adet=1,
            ants_is_urun_kodu="OTHER",
            aciklama="B",
            birim_fiyat=Decimal("1"),
            toplam_fiyat=Decimal("1"),
            stok_kodu="999999",
        ),
    ]
    teslim = "ANT5307 - 3 adet T0+30 Hafta\nT0: Sipariş Onay Tarihi"
    sync_quote_lines_to_import_excel(path, lines, teslim_tarihi_text=teslim)
    wb = load_workbook(path, data_only=True)
    assert wb.active.cell(3, 8).value == 220
    wb.close()


def test_sync_same_antsis_code_two_rows_keep_distinct_temin(tmp_path):
    """Aynı ürün kodu 2 satır: 5x12 ve 5x26 → her satır kendi Temin'ini alır."""
    path = tmp_path / "dup.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["meta"])
    ws.append(
        [
            "Teklifte Bulun",
            "Stok Kodu",
            "Stok Tanımı",
            "Miktar",
            "Fiyat",
            "Satır Toplamı",
            "Tedarikçi Notu",
            "Temin Süresi (Takvim Günü)",
            "Garanti Süresi (Yıl)",
        ]
    )
    ws.append(["Y", "111", "A1", 5, None, None, None, None, None])
    ws.append(["Y", "222", "A2", 5, None, None, None, None, None])
    wb.save(path)
    wb.close()

    code = "ANT-DUP"
    lines = [
        QuoteLine(
            row_number=1,
            adet=5,
            ants_is_urun_kodu=code,
            aciklama="Birinci",
            birim_fiyat=Decimal("10"),
            toplam_fiyat=Decimal("50"),
            stok_kodu="111",
            teslim_suresi_parcasi=f"{code} - 5 adet T0+12 Hafta",
        ),
        QuoteLine(
            row_number=2,
            adet=5,
            ants_is_urun_kodu=code,
            aciklama="Ikinci",
            birim_fiyat=Decimal("10"),
            toplam_fiyat=Decimal("50"),
            stok_kodu="222",
            teslim_suresi_parcasi=f"{code} - 5 adet T0+26 Hafta",
        ),
    ]
    teslim = (
        f"{code} - 5 adet T0+12 Hafta\n"
        f"{code} - 5 adet T0+26 Hafta\n"
        "T0: Sipariş Onay Tarihi"
    )
    sync_quote_lines_to_import_excel(path, lines, teslim_tarihi_text=teslim)
    out = load_workbook(path, data_only=True).active
    # 12 hafta → 84 → 90; 26 hafta → 182 → 190
    assert out.cell(3, 8).value == 90
    assert out.cell(4, 8).value == 190
    assert "T0+12" in str(out.cell(3, 7).value or "")
    assert "T0+26" in str(out.cell(4, 7).value or "")
