from __future__ import annotations

import re
import shutil
import tempfile
import uuid
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from teklif_hazirlama.core.models import SheetsRow
from teklif_hazirlama.core.sevk_tarihi_parser import (
    build_tedarikci_notu,
    parse_sevk_tarihi,
)
from teklif_hazirlama.infrastructure.excel_import_reader import (
    FIELD_ALIASES,
    HEADER_ROW,
    DATA_START_ROW,
    MAX_HEADER_COLS,
    ImportSheetReader,
    resolve_field_columns,
)

NA = "na"


class ExcelEnrichError(Exception):
    pass


@dataclass
class EnrichResult:
    enriched_path: Path
    matched: int = 0  # doldurulan satır sayısı
    warnings: list[str] = field(default_factory=list)


def _is_na(value: str | None) -> bool:
    return not (value or "").strip() or (value or "").strip().lower() == NA


def _parse_price(value: str) -> Decimal | None:
    if _is_na(value):
        return None
    text = re.sub(r"[A-Za-z$€₺]", "", (value or "").strip()).replace(" ", "")
    if text.count(",") == 1 and (text.count(".") >= 1 or text.count(",") == 1):
        if "." in text and "," in text:
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", ".")
    else:
        text = text.replace(",", ".")
    try:
        return Decimal(text)
    except (InvalidOperation, ValueError):
        return None


def _fill_values_from_sheet(sheet: SheetsRow) -> dict[str, Any]:
    """Sheets satırından Excel'e yazılacak alanlar (eksik/na → 'na')."""
    out: dict[str, Any] = {}

    price = _parse_price(sheet.birim_fiyat)
    out["fiyat"] = float(price) if price is not None else NA

    sevk = parse_sevk_tarihi(sheet.teklif_sevk_tarihi)
    if sevk.temin_gun is not None:
        out["temin_suresi"] = sevk.temin_gun
    else:
        out["temin_suresi"] = NA

    code = (sheet.antsis_urun_kodu or "").strip()
    if _is_na(code) or not sevk.entries:
        out["tedarikci_notu"] = NA
    else:
        ozel = sheet.kalemdeki_ozel_sartlar
        if _is_na(ozel):
            ozel = ""
        out["tedarikci_notu"] = build_tedarikci_notu(code, sevk.entries, ozel)

    out["garanti_suresi"] = 1.0
    return out


def _apply_fill_to_cells(
    cells: dict[str, Any], fill: dict[str, Any], columns: dict[str, int]
) -> dict[str, Any]:
    out = dict(cells)
    for field, value in fill.items():
        if field in columns:
            out[field] = value
    return out


def _unique_dest(dest_dir: Path, source: Path) -> Path:
    """Her seferinde yeni dosya — Windows dosya kilidi (WinError 32) önlenir."""
    dest_dir.mkdir(parents=True, exist_ok=True)
    return dest_dir / f"{source.stem}_dolu_{uuid.uuid4().hex[:8]}{source.suffix}"


class ExcelEnricher:
    """
    Sheets satırlarını Excel veri satırlarına **sırayla** yazar.
    Stok kodu eşleştirmesi yok. Boş/na → Excel'e 'na'.
    """

    def enrich(
        self,
        excel_path: str | Path,
        sheets_rows: list[SheetsRow],
        *,
        output_dir: str | Path | None = None,
    ) -> EnrichResult:
        source = Path(excel_path)
        if not source.exists():
            raise ExcelEnrichError(f"Excel bulunamadı: {source}")
        if not sheets_rows:
            raise ExcelEnrichError(
                "Sheets satırı yok; önce Google satırlarını yapıştırın."
            )

        dest_dir = (
            Path(output_dir)
            if output_dir
            else Path(tempfile.mkdtemp(prefix="teklif_enrich_"))
        )
        dest = _unique_dest(dest_dir, source)

        try:
            return self._enrich_openpyxl_copy(source, dest, sheets_rows)
        except ExcelEnrichError:
            raise
        except Exception as exc:
            dest2 = _unique_dest(dest_dir, source)
            result = self._enrich_rebuild_clean(source, dest2, sheets_rows)
            result.warnings.insert(
                0,
                "ERP Excel openpyxl ile açılamadı (bozuk XML); "
                f"temiz kopya üretildi. ({exc})",
            )
            return result

    def _enrich_openpyxl_copy(
        self,
        source: Path,
        dest: Path,
        sheets_rows: list[SheetsRow],
    ) -> EnrichResult:
        shutil.copy2(source, dest)
        warnings: list[str] = []
        wb = None
        filled = 0
        try:
            wb = load_workbook(dest)
            ws = wb.active
            headers = [
                ws.cell(row=HEADER_ROW, column=col).value
                for col in range(1, MAX_HEADER_COLS + 1)
            ]
            columns = resolve_field_columns(headers, aliases=FIELD_ALIASES)
            self._ensure_required_columns(columns, warnings)

            data_rows = self._list_excel_data_rows(ws, columns)
            filled = self._pair_and_warn(len(data_rows), len(sheets_rows), warnings)

            for i in range(filled):
                row_idx = data_rows[i]
                fill = _fill_values_from_sheet(sheets_rows[i])
                for field, value in fill.items():
                    if field not in columns:
                        continue
                    ws.cell(row=row_idx, column=columns[field] + 1).value = value

            for i in range(filled, len(data_rows)):
                row_idx = data_rows[i]
                for field, value in (
                    ("fiyat", NA),
                    ("temin_suresi", NA),
                    ("tedarikci_notu", NA),
                    ("garanti_suresi", 1.0),
                ):
                    if field in columns:
                        ws.cell(
                            row=row_idx, column=columns[field] + 1
                        ).value = value

            if filled == 0 and not data_rows:
                raise ExcelEnrichError("Excel'de doldurulacak veri satırı yok.")
            wb.save(dest)
        finally:
            if wb is not None:
                try:
                    wb.close()
                except Exception:
                    pass

        return EnrichResult(
            enriched_path=dest, matched=filled, warnings=warnings
        )

    def _enrich_rebuild_clean(
        self,
        source: Path,
        dest: Path,
        sheets_rows: list[SheetsRow],
    ) -> EnrichResult:
        reader = ImportSheetReader()
        warnings: list[str] = []
        try:
            headers = reader.header_row_values(source, header_row=HEADER_ROW)
        except Exception as exc:
            raise ExcelEnrichError(
                f"ERP Excel okunamadı: {source.name}. ({exc})"
            ) from exc

        header_list = [
            ("" if h is None else str(h)) for h in headers[:MAX_HEADER_COLS]
        ]
        while len(header_list) < MAX_HEADER_COLS:
            header_list.append("")

        columns = resolve_field_columns(header_list, aliases=FIELD_ALIASES)
        header_list = self._ensure_headers_for_enrich(header_list, columns, warnings)
        columns = resolve_field_columns(header_list, aliases=FIELD_ALIASES)
        self._ensure_required_columns(columns, warnings)

        try:
            raw_rows = reader.read_rows(
                source, header_row=HEADER_ROW, data_start_row=DATA_START_ROW
            )
        except Exception as exc:
            raise ExcelEnrichError(
                f"ERP Excel satırları okunamadı: {source.name}. ({exc})"
            ) from exc

        filled = self._pair_and_warn(len(raw_rows), len(sheets_rows), warnings)
        enriched_rows: list[dict[str, Any]] = []
        for i, (_row_idx, cells) in enumerate(raw_rows):
            if i < filled:
                fill = _fill_values_from_sheet(sheets_rows[i])
            else:
                fill = {
                    "fiyat": NA,
                    "temin_suresi": NA,
                    "tedarikci_notu": NA,
                    "garanti_suresi": 1.0,
                }
            enriched_rows.append(_apply_fill_to_cells(cells, fill, columns))

        if not enriched_rows:
            raise ExcelEnrichError("Excel'de doldurulacak veri satırı yok.")

        wb = Workbook()
        try:
            ws = wb.active
            ws.title = "Sayfa1"
            ws.append(["N"])
            ws.append(header_list)
            col_count = max(
                (i for i, h in enumerate(header_list) if str(h).strip()),
                default=0,
            ) + 1
            for cells in enriched_rows:
                values = [None] * col_count
                for field, idx in columns.items():
                    if idx < col_count:
                        values[idx] = cells.get(field)
                ws.append(values)
            wb.save(dest)
        finally:
            try:
                wb.close()
            except Exception:
                pass

        return EnrichResult(
            enriched_path=dest, matched=filled, warnings=warnings
        )

    def _list_excel_data_rows(self, ws, columns: dict[str, int]) -> list[int]:
        """Boş olmayan veri satırı indeksleri (sırayla)."""
        rows: list[int] = []
        max_row = ws.max_row or DATA_START_ROW
        stok_col = columns.get("stok_kodu")
        for row_idx in range(DATA_START_ROW, max_row + 1):
            if stok_col is not None:
                stok = ws.cell(row=row_idx, column=stok_col + 1).value
                if stok is None or str(stok).strip() == "":
                    empty = True
                    for col in range(1, MAX_HEADER_COLS + 1):
                        v = ws.cell(row=row_idx, column=col).value
                        if v is not None and str(v).strip():
                            empty = False
                            break
                    if empty:
                        continue
            rows.append(row_idx)
        return rows

    def _pair_and_warn(
        self, excel_count: int, sheets_count: int, warnings: list[str]
    ) -> int:
        filled = min(excel_count, sheets_count)
        if sheets_count > excel_count:
            warnings.append(
                f"Sheets'te {sheets_count - excel_count} fazla satır var; "
                "Excel satır sayısına göre sırayla dolduruldu."
            )
        if excel_count > sheets_count:
            warnings.append(
                f"Excel'de {excel_count - sheets_count} fazla satır var; "
                "eksik Sheets karşılıkları 'na' yazıldı."
            )
        return filled

    def _ensure_required_columns(
        self, columns: dict[str, int], warnings: list[str]
    ) -> None:
        missing = [
            f for f in ("stok_kodu", "fiyat", "tedarikci_notu") if f not in columns
        ]
        if missing:
            labels = [FIELD_ALIASES[f][0] for f in missing]
            raise ExcelEnrichError(
                "Excel'de zenginleştirme için eksik sütun(lar): "
                + ", ".join(labels)
            )
        if "temin_suresi" not in columns:
            warnings.append(
                "Excel'de 'Temin Süresi (Takvim Günü)' sütunu yok; Temin yazılamadı."
            )
        if "garanti_suresi" not in columns:
            warnings.append(
                "Excel'de 'Garanti Süresi (Yıl)' sütunu yok; Garanti yazılamadı."
            )

    def _ensure_headers_for_enrich(
        self,
        header_list: list[str],
        columns: dict[str, int],
        warnings: list[str],
    ) -> list[str]:
        out = list(header_list)
        extras = [
            ("temin_suresi", "Temin Süresi (Takvim Günü)"),
            ("garanti_suresi", "Garanti Süresi (Yıl)"),
        ]
        for field, title in extras:
            if field in columns:
                continue
            placed = False
            for i, h in enumerate(out):
                if not str(h).strip():
                    out[i] = title
                    placed = True
                    break
            if not placed:
                out.append(title)
            warnings.append(f"Eksik sütun eklendi: {title}")
        return out
