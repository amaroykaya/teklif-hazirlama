"""Arayüz ürün satırlarını doldurulmuş Import Excel'e yazar."""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from teklif_hazirlama.core.models import QuoteLine
from teklif_hazirlama.core.sevk_tarihi_parser import (
    build_tedarikci_notu,
    compute_temin_gun,
    resolve_teslim_entries_for_line,
)
from teklif_hazirlama.core.tedarikci_notu_parser import parse_tedarikci_notu
from teklif_hazirlama.infrastructure.excel_import_reader import (
    DATA_START_ROW,
    HEADER_ROW,
    MAX_HEADER_COLS,
    resolve_field_columns,
)


def sync_quote_lines_to_import_excel(
    excel_path: str | Path,
    lines: list[QuoteLine],
    teslim_tarihi_text: str = "",
) -> list[str]:
    """
    Teklifte Bulun = Y satırlarını UI satırlarıyla sırayla eşleştirir.
    Fazla UI satırı ekler; fazla Excel satırını siler.
    Teslim Tarihi metninden Temin Süresi (ve Tedarikçi Notu dilimleri) güncellenir.
    """
    path = Path(excel_path)
    if not path.exists():
        raise FileNotFoundError(f"Excel bulunamadı: {path}")
    if not lines:
        return ["Ürün satırı yok; doldurulmuş Excel güncellenmedi."]

    wb = load_workbook(path)
    try:
        ws = wb.active
        headers = [
            ws.cell(row=HEADER_ROW, column=col).value
            for col in range(1, MAX_HEADER_COLS + 1)
        ]
        columns = resolve_field_columns(headers)
        missing = [f for f in ("miktar", "fiyat") if f not in columns]
        if missing:
            return [
                "Doldurulmuş Excel güncellenemedi; eksik sütun: "
                + ", ".join(missing)
            ]

        warnings: list[str] = []
        data_rows = _list_included_rows(ws, columns)
        before = len(data_rows)
        teslim_text = (teslim_tarihi_text or "").strip()

        if before > len(lines):
            for row_idx in reversed(data_rows[len(lines) :]):
                ws.delete_rows(row_idx, 1)
            warnings.append(
                f"{before - len(lines)} fazla Excel satırı silindi (arayüzde yok)."
            )
            data_rows = _list_included_rows(ws, columns)

        temin_updates = 0
        for i, line in enumerate(lines):
            if i < len(data_rows):
                if _write_line(
                    ws, data_rows[i], columns, line, teslim_text, line_index=i
                ):
                    temin_updates += 1
            else:
                new_row, temin_ok = _append_row(
                    ws, columns, line, teslim_text, line_index=i
                )
                data_rows.append(new_row)
                if temin_ok:
                    temin_updates += 1

        added = max(0, len(lines) - before)
        if added:
            warnings.append(f"{added} yeni satır Excel'e eklendi.")
        if teslim_text and temin_updates:
            warnings.append(
                f"Temin Süresi, Teslim Tarihi alanından güncellendi ({temin_updates} satır)."
            )
        elif teslim_text:
            warnings.append(
                "Teslim Tarihi'nde T0+hafta dilimi okunamadı; Temin eski kaldı."
            )

        wb.save(path)
        warnings.insert(
            0,
            f"Doldurulmuş Excel arayüz satırlarıyla güncellendi ({len(lines)} satır).",
        )
        return warnings
    finally:
        wb.close()


def _list_included_rows(ws, columns: dict[str, int]) -> list[int]:
    rows: list[int] = []
    max_row = ws.max_row or DATA_START_ROW
    teklifte_col = columns.get("teklifte_bulun")
    stok_col = columns.get("stok_kodu")
    for row_idx in range(DATA_START_ROW, max_row + 1):
        if teklifte_col is not None:
            if not _is_teklifte_y(ws.cell(row=row_idx, column=teklifte_col + 1).value):
                continue
        elif stok_col is not None:
            stok = ws.cell(row=row_idx, column=stok_col + 1).value
            if stok is None or str(stok).strip() == "":
                if _row_empty(ws, row_idx):
                    continue
        else:
            if _row_empty(ws, row_idx):
                continue
        rows.append(row_idx)
    return rows


def _is_teklifte_y(value) -> bool:
    if value is True:
        return True
    text = str(value or "").strip().upper()
    return text in {"Y", "YES", "EVET", "1"}


def _row_empty(ws, row_idx: int) -> bool:
    for col in range(1, MAX_HEADER_COLS + 1):
        v = ws.cell(row=row_idx, column=col).value
        if v is not None and str(v).strip():
            return False
    return True


def _write_line(
    ws,
    row_idx: int,
    columns: dict[str, int],
    line: QuoteLine,
    teslim_tarihi_text: str = "",
    *,
    line_index: int = 0,
) -> bool:
    """True dönerse Temin/Teslim dilimleri UI metninden güncellendi."""
    _set(ws, row_idx, columns, "miktar", int(line.adet))
    _set(ws, row_idx, columns, "fiyat", _to_number(line.birim_fiyat))
    if "satir_toplami" in columns:
        toplam = line.toplam_fiyat
        if toplam == 0 and line.adet:
            toplam = line.birim_fiyat * line.adet
        _set(ws, row_idx, columns, "satir_toplami", _to_number(toplam))

    stok_kodu = (line.stok_kodu or "").strip()
    stok_tanimi = _stok_tanimi_from_line(line)
    if stok_kodu and "stok_kodu" in columns:
        _set(ws, row_idx, columns, "stok_kodu", stok_kodu)
    if stok_tanimi and "stok_tanimi" in columns:
        _set(ws, row_idx, columns, "stok_tanimi", stok_tanimi)

    code = (line.ants_is_urun_kodu or "").strip()
    existing_notu = None
    if "tedarikci_notu" in columns:
        existing_notu = ws.cell(
            row=row_idx, column=columns["tedarikci_notu"] + 1
        ).value

    teslim_text = (teslim_tarihi_text or "").strip()
    entries = []
    if teslim_text:
        # UI Teslim Tarihi öncelikli — eski import parçasını KULLANMA
        entries = resolve_teslim_entries_for_line(
            teslim_text, code, line_index=line_index
        )
    elif line.teslim_suresi_parcasi:
        entries = resolve_teslim_entries_for_line(
            f"{code} - {line.teslim_suresi_parcasi}",
            code,
            line_index=0,
        )

    if entries:
        ozel = ""
        if existing_notu:
            ozel = parse_tedarikci_notu(str(existing_notu)).aciklama_eki
        if "tedarikci_notu" in columns:
            _set(
                ws,
                row_idx,
                columns,
                "tedarikci_notu",
                build_tedarikci_notu(code or "na", entries, ozel),
            )
        temin = compute_temin_gun(entries)
        if temin is not None and "temin_suresi" in columns:
            _set(ws, row_idx, columns, "temin_suresi", temin)
        return True

    if "tedarikci_notu" in columns:
        _set(
            ws,
            row_idx,
            columns,
            "tedarikci_notu",
            _update_tedarikci_notu_code(existing_notu, code),
        )
    return False


def _append_row(
    ws,
    columns: dict[str, int],
    line: QuoteLine,
    teslim_tarihi_text: str = "",
    *,
    line_index: int = 0,
) -> tuple[int, bool]:
    new_row = (ws.max_row or DATA_START_ROW - 1) + 1
    if "teklifte_bulun" in columns:
        _set(ws, new_row, columns, "teklifte_bulun", "Y")
    stok_kodu = (line.stok_kodu or "").strip() or "manuel"
    if "stok_kodu" in columns:
        _set(ws, new_row, columns, "stok_kodu", stok_kodu)
    if "stok_tanimi" in columns:
        _set(
            ws,
            new_row,
            columns,
            "stok_tanimi",
            _stok_tanimi_from_line(line) or line.aciklama,
        )
    temin_ok = _write_line(
        ws, new_row, columns, line, teslim_tarihi_text, line_index=line_index
    )
    if "garanti_suresi" in columns:
        cell = ws.cell(row=new_row, column=columns["garanti_suresi"] + 1)
        if cell.value is None or str(cell.value).strip() == "":
            cell.value = 1
    if "temin_suresi" in columns:
        cell = ws.cell(row=new_row, column=columns["temin_suresi"] + 1)
        if cell.value is None or str(cell.value).strip() == "":
            cell.value = "na"
    return new_row, temin_ok


def _set(ws, row_idx: int, columns: dict[str, int], field: str, value) -> None:
    if field not in columns:
        return
    ws.cell(row=row_idx, column=columns[field] + 1).value = value


def _to_number(value: Decimal) -> float:
    return float(value)


def _stok_tanimi_from_line(line: QuoteLine) -> str:
    text = (line.aciklama or "").strip()
    if not text:
        return (line.stok_aciklama or "").strip()
    first = text.split("\n", 1)[0].strip()
    stok_kodu = (line.stok_kodu or "").strip()
    if stok_kodu and first.endswith(f" / {stok_kodu}"):
        return first[: -len(f" / {stok_kodu}")].strip()
    if " / " in first and stok_kodu:
        left, right = first.rsplit(" / ", 1)
        if right.strip() == stok_kodu:
            return left.strip()
    return first


def _update_tedarikci_notu_code(existing, new_code: str) -> str:
    code = (new_code or "").strip()
    raw = "" if existing is None else str(existing).strip()
    if not code or code.lower() == "na":
        return raw if raw else "na"
    if not raw or raw.lower() == "na":
        return f"Teslim süresi : {code} - T0: Sipariş Onay Tarihi"

    parsed = parse_tedarikci_notu(raw)
    if not parsed.urun_kodu:
        return f"Teslim süresi : {code} - T0: Sipariş Onay Tarihi"

    parcasi = parsed.teslim_suresi_parcasi or ""
    if " - " in parcasi:
        rest = parcasi.split(" - ", 1)[1]
        segment = f"{code} - {rest}"
    else:
        segment = f"{code} - {parcasi}".strip(" -")

    out = f"Teslim süresi : {segment} T0: Sipariş Onay Tarihi"
    if parsed.aciklama_eki:
        out = f"{out} {parsed.aciklama_eki}"
    return out
