from __future__ import annotations

import math
import shutil
from copy import copy
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from teklif_hazirlama.core.models import QuoteDocument
from teklif_hazirlama.core.quote_engine import QuoteEngine
from teklif_hazirlama.paths import REVIZE_TEMPLATE_FILE, TEMPLATE_FILE

DEFAULT_COLUMN_WIDTH = 8.43
DEFAULT_FONT_SIZE = 11.0
LINE_HEIGHT_FACTOR = 1.45
# Excel genişlik birimi gerçek karakter sığmasından iyimser; güven payı bırak.
WIDTH_CHAR_FACTOR = 0.72
ROW_HEIGHT_PADDING = 6.0
DATE_NUMBER_FORMAT = "DD.MM.YYYY"


class TemplateFiller:
    PRODUCT_FIRST_ROW = 27
    TEMPLATE_PRODUCT_ROWS = 2
    TESLIM_TARIHI_ROW = 23
    TESLIM_TARIHI_COLUMNS = (2, 3, 4)
    TESLIMAT_COLUMNS = (5, 6)  # E Teslimat, F Teslimat Şekli
    ODEME_COLUMNS = (7, 8)

    def __init__(self, *, revise: bool = False) -> None:
        self.revise = revise
        if revise:
            self.sheet_name = "Revize Teklif"
            self.template_file = REVIZE_TEMPLATE_FILE
            # Açıklama C:D birleşik; E Birim, F İndirim, G İndirimli Birim, H İndirimli Toplam
            self.aciklama_columns = (3, 4)
            self.aciklama_merge_end = 4
        else:
            self.sheet_name = "Teklif"
            self.template_file = TEMPLATE_FILE
            self.aciklama_columns = (3, 4, 5, 6)
            self.aciklama_merge_end = 6

    def fill(self, document: QuoteDocument, output_path: str | Path) -> Path:
        output = Path(output_path)
        output.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(self.template_file, output)

        wb = load_workbook(output)
        ws = wb[self.sheet_name]

        n = len(document.lines)
        extra = max(0, n - self.TEMPLATE_PRODUCT_ROWS)
        unused = max(0, self.TEMPLATE_PRODUCT_ROWS - n) if n else 0
        if extra:
            ws.insert_rows(self.PRODUCT_FIRST_ROW + self.TEMPLATE_PRODUCT_ROWS, extra)
            for i in range(extra):
                row = self.PRODUCT_FIRST_ROW + self.TEMPLATE_PRODUCT_ROWS + i
                self._copy_row_style(ws, self.PRODUCT_FIRST_ROW, row)
                ws.merge_cells(
                    start_row=row,
                    start_column=3,
                    end_row=row,
                    end_column=self.aciklama_merge_end,
                )
                if self.revise:
                    ws[f"F{row}"] = f"=1-(G{row}/E{row})"
                    ws[f"H{row}"] = f"=A{row}*G{row}"
        elif unused:
            # Şablon 2 ürün satırlı; tek satırda boş satır + eski Toplam formülü (#DEĞER!) kalmasın.
            ws.delete_rows(self.PRODUCT_FIRST_ROW + n, unused)

        summary_row = self.PRODUCT_FIRST_ROW + n
        customer = document.customer
        header = document.header
        engine = QuoteEngine()

        ws["H8"] = header.tarih
        ws["H8"].number_format = DATE_NUMBER_FORMAT
        ws["H9"] = header.teklif_no
        ws["H10"] = engine.gecerlilik_tarihi(customer, header.tarih)
        ws["H10"].number_format = DATE_NUMBER_FORMAT
        ws["H11"] = header.hazirlayan
        self._fit_teklif_no_column(ws, header.teklif_no)
        ws["A14"] = f"Sayın {header.hitap_kisi}"
        ws["A16"] = customer.firma.get("unvan", customer.name)
        adres = customer.firma.get("adres_satirlari", [])
        telefon = header.hitap_telefon or customer.firma.get("telefon", "")
        for idx, row_num in enumerate([17, 18, 19]):
            if idx < len(adres):
                ws[f"A{row_num}"] = adres[idx]
            else:
                ws[f"A{row_num}"] = ""
        if telefon:
            ws["A19"] = telefon

        ws["A23"] = header.istek_no
        ws["B23"] = document.teslim_tarihi_text
        ws["E23"] = header.teslimat
        ws["F23"] = header.teslimat_sekli
        ws["G23"] = header.odeme_sekli
        self._fit_column_to_text(ws, 5, header.teslimat, max_width=22)
        self._fit_column_to_text(ws, 6, header.teslimat_sekli, max_width=24)
        self._apply_wrap_text(ws, self.TESLIM_TARIHI_ROW, self.TESLIM_TARIHI_COLUMNS)
        self._apply_wrap_text(ws, self.TESLIM_TARIHI_ROW, self.TESLIMAT_COLUMNS)
        self._apply_wrap_text(ws, self.TESLIM_TARIHI_ROW, self.ODEME_COLUMNS)
        self._autofit_row_height(
            ws,
            self.TESLIM_TARIHI_ROW,
            [
                self.TESLIM_TARIHI_COLUMNS,
                self.TESLIMAT_COLUMNS,
                self.ODEME_COLUMNS,
            ],
        )

        currency = customer.para_birimi
        for i, line in enumerate(document.lines):
            row = self.PRODUCT_FIRST_ROW + i
            ws[f"A{row}"] = line.adet
            ws[f"B{row}"] = line.ants_is_urun_kodu
            ws[f"C{row}"] = line.aciklama
            if self.revise:
                # E Birim Fiyat, G İndirimli Birim Fiyat; F = 1-(G/E)
                birim = float(line.birim_fiyat)
                indirimli = float(
                    line.indirimli_birim_fiyat
                    if line.indirimli_birim_fiyat is not None
                    else line.birim_fiyat
                )
                ws[f"E{row}"] = birim
                ws[f"G{row}"] = indirimli
                ws[f"F{row}"] = f"=1-(G{row}/E{row})"
                ws[f"H{row}"] = format_money(line.toplam_fiyat, currency)
            else:
                ws[f"G{row}"] = format_money(line.birim_fiyat, currency)
                ws[f"H{row}"] = format_money(line.toplam_fiyat, currency)
            self._apply_wrap_text(
                ws, row, self.aciklama_columns, vertical="top"
            )
            self._autofit_row_height(ws, row, [self.aciklama_columns])

        ws[f"G{summary_row}"] = "Ara Toplam"
        ws[f"H{summary_row}"] = format_money(document.ara_toplam, currency)
        ws[f"G{summary_row + 1}"] = "KDV (%20)"
        ws[f"H{summary_row + 1}"] = format_money(document.kdv, currency)
        ws[f"G{summary_row + 2}"] = "Toplam"
        ws[f"H{summary_row + 2}"] = format_money(document.toplam, currency)

        self._write_sartlar_and_ozel(ws, document, summary_row)
        self._clear_stray_merges_in_sartlar(ws, summary_row)
        last_content_row = self._last_content_row(ws, summary_row)
        self._update_print_area(ws, last_content_row)
        self._make_workbook_editable(wb)

        wb.save(output)
        wb.close()
        self._unblock_windows_file(output)
        return output

    def _make_workbook_editable(self, wb) -> None:
        """Çıktı korumasız ve düzenlenebilir olsun (temiz Excel)."""
        from openpyxl.styles import Protection

        # Yardımcı gizli sayfaları çıkar (ör. adres)
        for name in list(wb.sheetnames):
            if name != self.sheet_name:
                del wb[name]

        ws = wb[self.sheet_name]
        try:
            ws.protection.disable()
        except Exception:
            ws.protection.sheet = False
            ws.protection.password = None

        if getattr(wb, "security", None) is not None:
            try:
                wb.security.workbookPassword = None
                wb.security.revisionsPassword = None
                wb.security.lockStructure = False
                wb.security.lockWindows = False
            except Exception:
                pass

        unlocked = Protection(locked=False, hidden=False)
        for row in ws.iter_rows():
            for cell in row:
                cell.protection = unlocked

        try:
            wb.properties.docSecurity = 0
        except Exception:
            pass

    @staticmethod
    def _unblock_windows_file(path: Path) -> None:
        """Mark of the Web / Protected View bandını kaldır."""
        try:
            zone = Path(str(path) + ":Zone.Identifier")
            if zone.exists():
                zone.unlink()
        except OSError:
            pass

    def _find_sartlar_row(self, ws, summary_row: int) -> int | None:
        for row in range(summary_row + 1, summary_row + 40):
            value = ws.cell(row, 1).value
            if value and str(value).strip().lower().startswith("şartlar"):
                return row
        return None

    def _write_sartlar_and_ozel(self, ws, document: QuoteDocument, summary_row: int) -> None:
        """
        Şablon: teşekkür → Şartlar:
        Şartlar: hemen altına 2 satır açılır:
          1) Antsis ürün kodları
          2) Özel şart
        Diğer şart metinleri bunların altında kalır.
        """
        sartlar_row = self._find_sartlar_row(ws, summary_row)
        if sartlar_row is None:
            return

        sartlar = [s.strip() for s in document.sartlar if s.strip()]
        codes = document.urun_kodlari_sartlar.strip()
        ozel = [s.strip() for s in document.ozel_sartlar if s.strip()]

        # Önce kodlar + özel şart(lar) (her biri ayrı satır), sonra diğer şart metinleri
        top: list[str] = []
        if codes:
            top.append(codes)
        top.extend(ozel)
        content_lines = top + sartlar

        existing = 0
        row = sartlar_row + 1
        max_scan = (ws.max_row or row) + 1
        while row <= max_scan:
            value = ws.cell(row, 1).value
            if value is None or str(value).strip() == "":
                break
            existing += 1
            row += 1

        needed = len(content_lines)
        style_source = sartlar_row + 1
        if needed > existing:
            insert_at = sartlar_row + 1 + existing
            extra = needed - existing
            ws.insert_rows(insert_at, extra)
            for i in range(extra):
                self._copy_row_style(ws, style_source, insert_at + i)
        elif needed < existing:
            for i in range(needed, existing):
                ws.cell(sartlar_row + 1 + i, 1).value = None

        # Şablondaki gibi: her şart bir satır, hücreye wrap/sığdırma yok
        for i, text in enumerate(content_lines):
            row = sartlar_row + 1 + i
            cell = ws.cell(row, 1)
            cell.value = text
            current = cell.alignment
            cell.alignment = Alignment(
                horizontal=current.horizontal if current else None,
                vertical=current.vertical if current else "center",
                wrap_text=False,
                shrink_to_fit=False,
            )
            if ws.row_dimensions[style_source].height:
                ws.row_dimensions[row].height = ws.row_dimensions[style_source].height

    def _clear_stray_merges_in_sartlar(self, ws, summary_row: int) -> None:
        """Remove leftover merges that cover şartlar text (e.g. template B47:F48)."""
        sartlar_row = self._find_sartlar_row(ws, summary_row)
        if sartlar_row is None:
            return
        to_unmerge = [
            str(merged)
            for merged in ws.merged_cells.ranges
            if merged.min_row >= sartlar_row and merged.min_col <= 6 and merged.max_col >= 2
        ]
        for ref in to_unmerge:
            ws.unmerge_cells(ref)

    def _last_content_row(self, ws, summary_row: int) -> int:
        last_row = summary_row + 2
        for row in range(summary_row + 1, summary_row + 80):
            for col in range(1, 9):
                value = ws.cell(row, col).value
                if value is not None and str(value).strip():
                    last_row = max(last_row, row)
                    break
        return last_row

    def _update_print_area(self, ws, last_row: int) -> None:
        """Excel tarafında 1×1 sayfa fit; PDF export aynı ayarı COM ile tekrarlar."""
        end_row = max(last_row + 1, 1)
        ws.print_area = f"A1:H{end_row}"
        # openpyxl: fitToPage + width/height; scale None → Zoom kapalı
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.page_setup.scale = None
        try:
            ws.page_setup.fitToPage = True
        except Exception:
            pass
        if ws.sheet_properties.pageSetUpPr is not None:
            ws.sheet_properties.pageSetUpPr.fitToPage = True

    def _fit_teklif_no_column(self, ws, teklif_no: str) -> None:
        """Uzun teklif no H hücresinde kalsın diye sütun genişliğini ayarla."""
        text = str(teklif_no or "").strip()
        cell = ws["H9"]
        current = cell.alignment
        cell.alignment = Alignment(
            horizontal=current.horizontal if current else "left",
            vertical=current.vertical if current else "center",
            wrap_text=False,
            shrink_to_fit=True,
        )
        self._fit_column_to_text(ws, 8, text, max_width=36, factor=1.15)

    def _fit_column_to_text(
        self,
        ws,
        col: int,
        text: str,
        *,
        max_width: float = 28,
        factor: float = 1.2,
    ) -> None:
        """Hücre metnine göre sütun genişliğini artır (küçültmez)."""
        value = str(text or "").strip()
        if not value:
            return
        letter = get_column_letter(col)
        base = self._column_width_chars(ws, col)
        longest = max(len(line) for line in value.splitlines() or [value])
        needed = max(base, longest * factor + 1.5)
        ws.column_dimensions[letter].width = min(needed, max_width)

    def _apply_wrap_text(
        self,
        ws,
        row: int,
        columns: tuple[int, ...],
        *,
        vertical: str | None = None,
    ) -> None:
        for col in columns:
            cell = ws.cell(row, col)
            current = cell.alignment
            cell.alignment = Alignment(
                horizontal=current.horizontal if current else None,
                vertical=(
                    vertical
                    if vertical is not None
                    else (current.vertical if current else "top")
                ),
                wrap_text=True,
            )

    def _autofit_row_height(self, ws, row: int, column_groups: list[tuple[int, ...]]) -> None:
        heights = [
            self._estimate_row_height(ws, row, columns)
            for columns in column_groups
        ]
        heights = [h for h in heights if h > 0]
        if heights:
            current = ws.row_dimensions[row].height or 0
            ws.row_dimensions[row].height = max(current, max(heights))

    def _estimate_row_height(self, ws, row: int, columns: tuple[int, ...]) -> float:
        cell = ws.cell(row, columns[0])
        text = cell.value
        if text is None or not str(text).strip():
            return 0.0

        width_chars = sum(self._column_width_chars(ws, col) for col in columns)
        width_chars = max(width_chars * WIDTH_CHAR_FACTOR, 8.0)
        font_size = cell.font.size if cell.font and cell.font.size else DEFAULT_FONT_SIZE
        line_height = font_size * LINE_HEIGHT_FACTOR

        total_lines = 0
        for paragraph in str(text).splitlines() or [""]:
            line_len = max(len(paragraph), 1)
            total_lines += max(1, math.ceil(line_len / width_chars))

        return total_lines * line_height + ROW_HEIGHT_PADDING

    def _column_width_chars(self, ws, col: int) -> float:
        letter = get_column_letter(col)
        width = ws.column_dimensions[letter].width
        return width if width else DEFAULT_COLUMN_WIDTH

    def _copy_row_style(self, ws, source_row: int, target_row: int) -> None:
        ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height
        for col in range(1, 13):
            src = ws.cell(source_row, col)
            dst = ws.cell(target_row, col)
            if src.has_style:
                dst.font = copy(src.font)
                dst.border = copy(src.border)
                dst.fill = copy(src.fill)
                dst.number_format = copy(src.number_format)
                dst.protection = copy(src.protection)
                dst.alignment = copy(src.alignment)


def load_default_sartlar(
    template_path: str | Path | None = None,
    *,
    revise: bool = False,
) -> list[str]:
    """Şablondaki 'Şartlar:' altındaki sabit açıklama satırlarını okur."""
    if template_path is not None:
        path = Path(template_path)
    elif revise:
        path = REVIZE_TEMPLATE_FILE
    else:
        path = TEMPLATE_FILE
    wb = load_workbook(path, data_only=True)
    try:
        preferred = "Revize Teklif" if revise else "Teklif"
        if preferred in wb.sheetnames:
            ws = wb[preferred]
        elif "Teklif" in wb.sheetnames:
            ws = wb["Teklif"]
        elif "Revize Teklif" in wb.sheetnames:
            ws = wb["Revize Teklif"]
        else:
            ws = wb.active
        sartlar_row = None
        for row in range(1, (ws.max_row or 1) + 1):
            value = ws.cell(row, 1).value
            if value and str(value).strip().lower().startswith("şartlar"):
                sartlar_row = row
                break
        if sartlar_row is None:
            return []
        lines: list[str] = []
        for row in range(sartlar_row + 1, (ws.max_row or sartlar_row) + 1):
            value = ws.cell(row, 1).value
            text = str(value).strip() if value is not None else ""
            if not text:
                break
            lines.append(text)
        return lines
    finally:
        wb.close()


def format_money(amount: Decimal, currency: str = "USD") -> str:
    q = amount.quantize(Decimal("0.01"))
    text = f"{q:.2f}"
    integer, _, frac = text.partition(".")
    int_with_sep = f"{int(integer):,}".replace(",", ".")
    if currency == "USD":
        return f"${int_with_sep},{frac}"
    return f"{int_with_sep},{frac} {currency}"


def output_basename(istek_no: str, teklif_no: str) -> str:
    safe_istek = sanitize_filename(istek_no)
    safe_teklif = sanitize_filename(teklif_no)
    return f"teklif-{safe_istek}-{safe_teklif}"


def sanitize_filename(value: str) -> str:
    forbidden = '<>:"/\\|?*'
    result = "".join(c if c not in forbidden else "-" for c in value.strip())
    return result or "cikti"
