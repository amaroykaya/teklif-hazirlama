from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

HEADER_ROW = 2
DATA_START_ROW = 3
MAX_HEADER_COLS = 40

# Alan anahtarı → 2. satırdaki normalize başlık ad(lar)ı (öncelik sırasıyla)
FIELD_ALIASES: dict[str, list[str]] = {
    "teklifte_bulun": ["teklifte bulun"],
    "stok_kodu": ["stok kodu"],
    "stok_tanimi": ["stok tanimi"],
    "miktar": ["miktar"],
    "termin_tarihi": ["termin tarihi", "gereksinim tarihi"],
    "fiyat": ["fiyat"],
    "satir_toplami": ["satir toplami"],
    "tedarikci_notu": ["tedarikci notu"],
    "kalite_provizyonlari": ["kalite provizyonlari"],
    "teknik_resim_sartname": [
        "teknik resim/sartname",
        "teknik resim sartname",
    ],
    "kalem_revizyon": ["kalem revizyon"],
    "temin_suresi": [
        "temin suresi (takvim gunu)",
        "temin suresi",
    ],
    "garanti_suresi": [
        "garanti suresi (yil)",
        "garanti suresi",
    ],
}

REQUIRED_FIELDS = (
    "teklifte_bulun",
    "stok_kodu",
    "stok_tanimi",
    "miktar",
    "fiyat",
    "satir_toplami",
    "tedarikci_notu",
)

# Teklif isteği aşamasında fiyat / tedarikçi notu henüz dolu olmayabilir
REQUIRED_FIELDS_ISTEK = (
    "teklifte_bulun",
    "stok_kodu",
    "stok_tanimi",
    "miktar",
)


def normalize_header(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    tr_map = str.maketrans(
        {
            "ı": "i",
            "İ": "i",
            "ğ": "g",
            "Ğ": "g",
            "ü": "u",
            "Ü": "u",
            "ş": "s",
            "Ş": "s",
            "ö": "o",
            "Ö": "o",
            "ç": "c",
            "Ç": "c",
        }
    )
    return " ".join(text.translate(tr_map).lower().split())


def resolve_field_columns(
    header_row: list[Any] | tuple[Any, ...],
    aliases: dict[str, list[str]] | None = None,
) -> dict[str, int]:
    """Başlık satırından alan → 0-based kolon indeksi haritası."""
    alias_map = aliases or FIELD_ALIASES
    header_to_index: dict[str, int] = {}
    for idx, value in enumerate(header_row):
        key = normalize_header(value)
        if key and key not in header_to_index:
            header_to_index[key] = idx

    resolved: dict[str, int] = {}
    for field, field_aliases in alias_map.items():
        for alias in field_aliases:
            if alias in header_to_index:
                resolved[field] = header_to_index[alias]
                break
    return resolved


class ImportSheetReader:
    """ERP export dosyalarini (data validation XML vb.) guvenle okur."""

    def read_rows(
        self,
        excel_path: str | Path,
        *,
        header_row: int = HEADER_ROW,
        data_start_row: int = DATA_START_ROW,
        aliases: dict[str, list[str]] | None = None,
    ) -> list[tuple[int, dict[str, Any]]]:
        path = Path(excel_path)
        columns = self.resolve_columns(
            path, header_row=header_row, aliases=aliases
        )
        errors: list[Exception] = []

        for loader in (self._load_calamine, self._load_read_only, self._load_openpyxl):
            try:
                return loader(
                    path,
                    columns,
                    data_start_row=data_start_row,
                )
            except Exception as exc:
                errors.append(exc)

        detail = " | ".join(str(e) for e in errors)
        raise ValueError(
            f"Import Excel okunamadı: {path.name}. "
            f"Dosya bozuk olabilir veya desteklenmeyen bir formattadır. ({detail})"
        )

    def resolve_columns(
        self,
        excel_path: str | Path,
        *,
        header_row: int = HEADER_ROW,
        aliases: dict[str, list[str]] | None = None,
    ) -> dict[str, int]:
        headers = self.header_row_values(excel_path, header_row=header_row)
        return resolve_field_columns(headers, aliases=aliases)

    def header_row_values(
        self, excel_path: str | Path, *, header_row: int = HEADER_ROW
    ) -> list[Any]:
        path = Path(excel_path)
        for loader in (self._header_row_calamine, self._header_row_read_only, self._header_row_openpyxl):
            try:
                return loader(path, header_row)
            except Exception:
                continue
        raise ValueError("Import Excel başlık satırı okunamadı")

    def _header_row_openpyxl(self, path: Path, header_row: int) -> list[Any]:
        wb = load_workbook(path, data_only=True)
        try:
            ws = wb.active
            return [
                ws.cell(row=header_row, column=col).value
                for col in range(1, MAX_HEADER_COLS + 1)
            ]
        finally:
            wb.close()

    def _header_row_read_only(self, path: Path, header_row: int) -> list[Any]:
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            row = next(
                wb.active.iter_rows(
                    min_row=header_row,
                    max_row=header_row,
                    min_col=1,
                    max_col=MAX_HEADER_COLS,
                    values_only=True,
                )
            )
            return list(row)
        finally:
            wb.close()

    def _header_row_calamine(self, path: Path, header_row: int) -> list[Any]:
        from python_calamine import CalamineWorkbook

        wb = CalamineWorkbook.from_path(str(path))
        matrix = wb.get_sheet_by_index(0).to_python()
        if len(matrix) < header_row:
            raise ValueError("Başlık satırı yok")
        row = list(matrix[header_row - 1])
        while len(row) < MAX_HEADER_COLS:
            row.append(None)
        return row[:MAX_HEADER_COLS]

    def _load_openpyxl(
        self, path: Path, columns: dict[str, int], *, data_start_row: int
    ) -> list[tuple[int, dict[str, Any]]]:
        wb = load_workbook(path, data_only=True)
        try:
            ws = wb.active
            rows: list[tuple[int, dict[str, Any]]] = []
            max_row = ws.max_row or data_start_row
            for row in range(data_start_row, max_row + 1):
                values = [
                    ws.cell(row=row, column=col).value
                    for col in range(1, MAX_HEADER_COLS + 1)
                ]
                data = self._dict_from_values(values, columns)
                if self._dict_empty(data):
                    continue
                rows.append((row, data))
            return rows
        finally:
            wb.close()

    def _load_read_only(
        self, path: Path, columns: dict[str, int], *, data_start_row: int
    ) -> list[tuple[int, dict[str, Any]]]:
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            rows: list[tuple[int, dict[str, Any]]] = []
            for offset, row_values in enumerate(
                wb.active.iter_rows(
                    min_row=data_start_row,
                    min_col=1,
                    max_col=MAX_HEADER_COLS,
                    values_only=True,
                ),
                start=data_start_row,
            ):
                if self._row_empty(row_values):
                    if offset > data_start_row and all(v is None for v in row_values):
                        break
                    continue
                rows.append((offset, self._dict_from_values(row_values, columns)))
            return rows
        finally:
            wb.close()

    def _load_calamine(
        self, path: Path, columns: dict[str, int], *, data_start_row: int
    ) -> list[tuple[int, dict[str, Any]]]:
        from python_calamine import CalamineWorkbook

        wb = CalamineWorkbook.from_path(str(path))
        matrix = wb.get_sheet_by_index(0).to_python()
        rows: list[tuple[int, dict[str, Any]]] = []
        for row_idx, values in enumerate(matrix, start=1):
            if row_idx < data_start_row:
                continue
            if self._row_empty(values):
                continue
            rows.append((row_idx, self._dict_from_values(values, columns)))
        return rows

    def _dict_from_values(self, values: tuple | list, columns: dict[str, int]) -> dict[str, Any]:
        data: dict[str, Any] = {}
        for field, idx in columns.items():
            data[field] = values[idx] if idx < len(values) else None
        return data

    def _row_empty(self, values: tuple | list) -> bool:
        return all(v is None or str(v).strip() == "" for v in values)

    def _dict_empty(self, data: dict[str, Any]) -> bool:
        return all(v is None or str(v).strip() == "" for v in data.values())
